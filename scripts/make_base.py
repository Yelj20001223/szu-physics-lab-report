"""Create the blank report base (assets/base.docx) from a finished report.

The school template only exists as a Word .doc for the students; every report
is a copy of it with the cover filled in. This tool takes any *finished*
report that still uses that template (for example the report you submitted
last week), blanks the cover values and the seven section cells, and writes a
clean starting document. Unused images/equations from the source are dropped.

Usage
-----
    python make_base.py --source finished_report.docx --out ../assets/base.docx
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_builder import SECTION_ORDER, fill_segment, underline_segments  # noqa: E402

COVER_LABELS = ("课程名称", "实验名称", "学院", "指导教师", "报告人", "学号",
                "实验时间", "提交时间")


def blank_cover(doc):
    for p in doc.paragraphs:
        stripped = p.text.replace(" ", "")
        if not stripped:
            continue
        if not any(stripped.startswith(label) for label in COVER_LABELS):
            continue
        for segment in underline_segments(p):
            text = "".join(r.text for r in segment)
            keep = "：" if text.strip().startswith("：") else ""
            fill_segment(segment, keep)
        # 封面行距保持模板原样（3.75 倍），report_builder 会在生成时按参考报告
        # 统一成固定 60 pt，使 Word 与其他阅读器里的行距一致


def tight_line_spacing(paragraph, line):
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = ppr.makeelement(qn("w:spacing"), {})
        ppr.append(spacing)
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def drop_spacer_paragraphs(doc):
    """Remove the empty filler paragraphs that sat between cover and body."""
    body = doc.element.body
    children = list(body)
    seen_cover_end = False
    for element in children:
        tag = element.tag.split("}")[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in element.iter(qn("w:t")))
            if text.strip().startswith("提交时间"):
                seen_cover_end = True
                continue
            if seen_cover_end and not text.strip():
                body.remove(element)
            elif seen_cover_end and text.strip():
                seen_cover_end = False


def blank_sections(doc):
    tables = [t for t in doc.tables if len(t.columns) == 1 and len(t.rows) >= 7]
    if not tables:
        raise SystemExit("no section table found in the source document")
    table = tables[0]
    for index, (key, _aliases, title) in enumerate(SECTION_ORDER):
        cell = table.rows[index].cells[0]
        for paragraph in cell.paragraphs[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        heading = cell.paragraphs[0]
        for run in list(heading.runs)[1:]:
            run._element.getparent().remove(run._element)
        if heading.runs:
            heading.runs[0].text = title
        else:
            heading.add_run(title)
        if not cell.paragraphs[0].runs:
            cell.paragraphs[0].add_run(title)
        # the source file numbered 一、二、三 through a Word list; the report
        # builder writes the number into the text instead, so drop the list
        ppr = heading._p.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is not None:
            ppr.remove(numpr)
    return table


def strip_unused_media(path):
    """Drop media parts that are no longer referenced, keeping the file small."""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        rels_map = {}
        for rel in [n for n in names if n.endswith(".rels")]:
            text = zf.read(rel).decode("utf-8", "ignore")
            owner = os.path.dirname(rel)
            owner = owner[:-5] if owner.endswith("_rels") else owner
            for rid, target in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', text):
                rels_map[(owner, rid)] = target
        used = set()
        for part in names:
            if not part.endswith(".xml") or part.endswith(".rels"):
                continue
            if not (part.startswith("word/") and "/media/" not in part):
                continue
            text = zf.read(part).decode("utf-8", "ignore")
            owner = os.path.dirname(part)
            for rid in re.findall(r'r:(?:embed|id|link|pict)="([^"]+)"', text):
                target = rels_map.get((owner, rid))
                if target:
                    used.add(os.path.basename(target))
        drop = [n for n in names
                if (n.startswith("word/media/") or n.startswith("word/embeddings/"))
                and os.path.basename(n) not in used]
        if not drop:
            return 0
        dropped_names = set(os.path.basename(n) for n in drop)

        def keep_or_drop(match):
            target = re.search(r'Target="([^"]+)"', match.group(0))
            if target and os.path.basename(target.group(1)) in dropped_names:
                return ""
            return match.group(0)

        tmp = path + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
            for item in zf.infolist():
                if item.filename in drop:
                    continue
                data = zf.read(item.filename)
                if item.filename.endswith(".rels"):
                    data = re.sub(r"<Relationship\b[^>]*/>", keep_or_drop,
                                  data.decode("utf-8")).encode("utf-8")
                out.writestr(item, data)
    shutil.move(tmp, path)
    return len(drop)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Create assets/base.docx")
    parser.add_argument("--source", required=True, help="a finished report .docx")
    parser.add_argument("--out", required=True, help="output base.docx")
    args = parser.parse_args(argv)

    doc = Document(args.source)
    blank_cover(doc)
    drop_spacer_paragraphs(doc)
    blank_sections(doc)
    out = os.path.abspath(args.out)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    dropped = strip_unused_media(out)
    print("WROTE %s (dropped %d unused media parts, %.0f KB)"
          % (out, dropped, os.path.getsize(out) / 1024.0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
