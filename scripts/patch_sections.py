#!/usr/bin/env python
"""把 content.md 里的指定小节，重新渲染进一份**已经存在的**报告 docx。

用在"用户在 Word 里改过报告，只想补/改某几节"的场景（用户 2026-10-13 要求）：
**不重新生成整份文档**，只替换正文表格里对应小节单元格的内容；封面、已贴的
预习试卷截图、手写原始数据照片、批阅/成绩表等**全部原样保留**。

    python patch_sections.py --docx "报告.docx" --content content-完整报告.md \
        --sections data,conclusion,questions

小节名（见 report_builder.SECTION_ORDER）：
    purpose principle apparatus procedure data conclusion questions

打完会自动打印**指纹对比**（封面段落、正文表格各行、内嵌图片 md5），用来确认
"只有指定小节变了、用户的其它改动没被碰"。默认**覆盖原文件**并留一份
`<报告.docx>.bak-<时间戳>` 备份；用 --out 另存、--no-backup 关掉备份。
"""
import argparse
import datetime
import hashlib
import os
import shutil
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx                                    # noqa: E402
from docx.oxml.ns import qn                    # noqa: E402

import report_builder as RB                    # noqa: E402

SECTION_TITLES = {k: title for (k, _aliases, title) in RB.SECTION_ORDER}
ROW_OF = {k: i for i, (k, _a, _t) in enumerate(RB.SECTION_ORDER)}  # purpose=0 … questions=6


def clear_nested_tables(cell):
    """删掉单元格里已有的嵌套表格（重复运行时保证幂等）。"""
    tc = cell._tc
    for tbl in tc.findall(qn("w:tbl")):
        tc.remove(tbl)


def _media_hashes(path):
    out = {}
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist()):
            if name.startswith("word/media/"):
                out[name] = hashlib.md5(z.read(name)).hexdigest()
    return out


def fingerprint(path):
    """封面段落 / 正文表格各行文本 / 内嵌图片 md5。"""
    d = docx.Document(path)
    fp = {
        "cover": [p.text for p in d.paragraphs],
        "rows": [(r.cells[0].text or "").strip() for r in d.tables[1].rows],
        "media": _media_hashes(path),
    }
    return fp


def report_diff(before, after, touched):
    print("--- 指纹对比（确认只有目标小节变了） ---")
    print("  封面段落：%s" % ("一致（未动）" if before["cover"] == after["cover"] else "!! 有变化 !!"))
    added = sorted(set(after["media"]) - set(before["media"]))
    removed = sorted(set(before["media"]) - set(after["media"]))
    changed = [n for n in before["media"]
               if n in after["media"] and before["media"][n] != after["media"][n]]
    print("  图片：%d 张 -> %d 张；新增 %s；删除 %s；内容变化 %s"
          % (len(before["media"]), len(after["media"]),
             added or "无", removed or "无", changed or "无"))
    for i, (b, a) in enumerate(zip(before["rows"], after["rows"])):
        flag = "SAME" if b == a else "CHANGED"
        note = "" if flag == "SAME" else (
            "  <= 目标小节" if i in [ROW_OF[k] for k in touched] else "  !! 非目标行 !!")
        print("  row%d [%s]%s" % (i, flag, note))


def main(argv=None):
    ap = argparse.ArgumentParser(description="把指定小节补进已有报告 docx")
    ap.add_argument("--docx", required=True, help="要打补丁的报告 docx")
    ap.add_argument("--content", required=True, help="content.md（小节内容源）")
    ap.add_argument("--sections", required=True,
                    help="逗号分隔的小节名，如 data,conclusion,questions")
    ap.add_argument("--out", default=None, help="输出路径，默认覆盖 --docx")
    ap.add_argument("--no-backup", action="store_true", help="不生成 .bak 备份")
    args = ap.parse_args(argv)

    keys = [s.strip() for s in args.sections.split(",") if s.strip()]
    bad = [k for k in keys if k not in SECTION_TITLES]
    if bad:
        raise SystemExit("未知小节 %s（可用：%s）" % (", ".join(bad), ", ".join(SECTION_TITLES)))
    if not keys:
        raise SystemExit("--sections 不能为空")

    content_path = os.path.abspath(args.content)
    base_dir = os.path.dirname(content_path)
    with open(content_path, encoding="utf-8") as fh:
        sections = RB.parse_content(fh.read())

    before = fingerprint(args.docx)

    doc = docx.Document(args.docx)
    if len(doc.tables) < 2 or len(doc.tables[1].rows) < 7:
        raise SystemExit("这个 docx 不像报告模板：找不到 7 行的正文表格")
    body = doc.tables[1]

    done = []
    for key in keys:
        blocks = sections.get(key, [])
        if not blocks:
            print("WARN 小节 %s 在 content 里是空的，跳过" % key, file=sys.stderr)
            continue
        cell = body.rows[ROW_OF[key]].cells[0]
        clear_nested_tables(cell)
        RB._render_cell(cell, SECTION_TITLES[key], blocks, base_dir, "full", key)
        done.append(key)

    out = os.path.abspath(args.out or args.docx)
    if not args.no_backup and out == os.path.abspath(args.docx):
        stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = args.docx + ".bak-" + stamp
        shutil.copy2(args.docx, backup)
        print("backup -> %s" % backup)
    doc.save(out)
    print("PATCHED %s : %s" % (out, ", ".join(done)))

    report_diff(before, fingerprint(out), done)


if __name__ == "__main__":
    main()
